from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def save_results_to_xls(results, filename="borda_results.xlsx", dir="results"):
    df = pd.DataFrame(results)
    df = df.round(4)

    output_dir = Path(__file__).resolve().parent.parent.parent / dir
    output_dir.mkdir(parents=True, exist_ok=True)  # create dir if not exists
    file_path = output_dir / filename

    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        sheet_name = 'Results'
        df.to_excel(writer, index=False, sheet_name=sheet_name)

        worksheet = writer.sheets[sheet_name]
        _apply_formatting(worksheet, df)

    print(f"--- Wyniki zapisano do pliku {file_path}")


def _apply_formatting(worksheet, df):
    full_range = worksheet.dimensions
    tab = Table(displayName="BordaResultsTable", ref=full_range)

    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showColumnStripes=False
    )

    worksheet.add_table(tab)

    # match width of columns
    for i, col in enumerate(df.columns, 1):
        max_data_len = df[col].astype(str).str.len().max()
        header_len = len(str(col))
        adjusted_width = (max(max_data_len, header_len) * 1.1) + 4
        worksheet.column_dimensions[get_column_letter(i)].width = adjusted_width